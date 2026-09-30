import json
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT

with open('scripts/inception_activity.json', encoding='utf-8') as f:
    r = json.load(f)
b = r['breakdowns']
inc = r['included']

doc = Document('Report Template - New.docx')
doc.paragraphs[0].add_run().add_break(WD_BREAK.PAGE)

title = doc.add_paragraph()
title.style = 'Normal'
title.add_run('GBMS Support App').bold = True
title.runs[0].font.size = Pt(22)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph('Inception-to-Date Activity Report')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].bold = True
doc.add_paragraph(
    f"Reporting period: {r['inception']['firstRetainedRecord'][:10]} to "
    f"{r['inception']['lastRetainedRecord'][:10]}\n"
    f"Prepared: {r['generatedAt'][:10]}\n"
    'Source: Firebase Firestore production activity records, project mofapp-60963'
)

doc.add_heading('Executive summary', 1)
doc.add_paragraph(
    f"This report covers all retained timestamped activity since the earliest available app record. "
    f"The data contains {len(inc['users']):,} user records, {len(inc['tickets']):,} operational tickets, "
    f"{len(inc['activities']):,} related ticket activities, {len(inc['notifications']):,} related notifications, "
    f"and {len(inc['auditLogs']):,} system audit-log events."
)
doc.add_paragraph(
    f"The test-ticket rule was applied consistently: tickets created by MMDA users on Tuesdays or Fridays "
    f"were excluded, while MDA-created tickets were retained. This removed {r['rules']['excludedTestTicketCount']} "
    'test tickets and their related ticket activities and ticket-linked notifications.'
)

def add_table(title, data, headers=('Item', 'Count')):
    doc.add_heading(title, 2)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for key, value in data.items():
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = f'{value:,}' if isinstance(value, int) else str(value)
    return table

add_table('Overall coverage', {
    'Users': len(inc['users']),
    'Operational tickets': len(inc['tickets']),
    'Ticket activities': len(inc['activities']),
    'Notifications': len(inc['notifications']),
    'System audit logs': len(inc['auditLogs']),
})
add_table('Users by role', b['usersByRole'])
add_table('Users by institution type', b['usersByInstitutionType'])
add_table('Users created by month', b['usersCreatedByMonth'])
add_table('Tickets by month', b['ticketsByMonth'])
add_table('Tickets by creator institution type', b['ticketsByInstitutionType'])
add_table('Tickets by creation weekday', b['ticketsByDay'])
add_table('Tickets by current status', b['ticketsByStatus'])
add_table('Tickets by priority', b['ticketsByPriority'])
add_table('Tickets by category', b['ticketsByCategory'])
add_table('Ticket activity actions', b['activitiesByAction'])
add_table('Notifications by type', b['notificationsByType'])
add_table('System audit actions', b['auditByAction'])
add_table('System audit events by month', b['auditByMonth'])

doc.add_heading('Operational ticket register', 1)
table = doc.add_table(rows=1, cols=6)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER
for cell, text in zip(table.rows[0].cells, ['Reference', 'Created', 'Institution', 'Category', 'Priority', 'Status']):
    cell.text = text
for ticket in sorted(inc['tickets'], key=lambda x: x.get('createdAt', '')):
    cells = table.add_row().cells
    values = [ticket.get('ticketReference', ticket['id']), ticket.get('createdAt', '')[:10], ticket.get('institutionId', ''), ticket.get('category', ''), ticket.get('priority', ''), ticket.get('status', '')]
    for cell, value in zip(cells, values):
        cell.text = str(value)

doc.add_heading('Excluded test tickets', 1)
doc.add_paragraph(
    f"{r['rules']['excludedTestTicketCount']} tickets were excluded under the MMDA Tuesday/Friday rule. "
    'Their full details are provided in excluded_test_tickets.csv.'
)

doc.add_heading('Findings', 1)
findings = [
    f"MMDA users created {b['ticketsByInstitutionType'].get('MMDA', 0)} operational tickets; MDA users created {b['ticketsByInstitutionType'].get('MDA', 0)}.",
    f"The most common operational category was {max(b['ticketsByCategory'], key=b['ticketsByCategory'].get)}.",
    f"The most frequent ticket activity was {max(b['activitiesByAction'], key=b['activitiesByAction'].get)}.",
    f"The audit trail contains {b['auditByAction'].get('user_created', 0):,} user-created events, {b['auditByAction'].get('user_updated', 0):,} user updates, {b['auditByAction'].get('password_reset', 0):,} password resets, and {b['auditByAction'].get('announcement_sent', 0):,} announcement event.",
]
for finding in findings:
    doc.add_paragraph('• ' + finding, style='Normal')

doc.add_heading('Data limitations', 1)
doc.add_paragraph(
    'This is a report of retained Firestore records. Deleted or purged documents cannot be reconstructed from the live database. '
    'Presence heartbeats overwrite the current user profile field and are not retained as historical sessions. Ordinary screen views '
    'and Firestore reads are not logged unless they create one of the covered records. Weekday classification uses the stored calendar date, '
    'which aligns with Ghana’s calendar date for the retained records.'
)

doc.add_heading('Complete appendices', 1)
doc.add_paragraph(
    'The accompanying CSV files contain the complete included ticket activity, audit logs, notifications, user-creation events, '
    'and excluded test-ticket list. The full structured export is inception_activity.json.'
)
doc.add_paragraph('Prepared for management submission.')

doc.save('GBMS Activity Report - Updated v4.docx')
print('Wrote GBMS Activity Report - Updated v4.docx')
