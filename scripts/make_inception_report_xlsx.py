import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent.parent
with (ROOT / 'scripts' / 'inception_activity.json').open(encoding='utf-8') as f:
    report = json.load(f)

wb = Workbook()
wb.remove(wb.active)

header_fill = PatternFill('solid', fgColor='1F4E78')
section_fill = PatternFill('solid', fgColor='D9EAF7')
header_font = Font(color='FFFFFF', bold=True)
title_font = Font(size=16, bold=True, color='1F4E78')


def flatten(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def add_table(name, rows, columns=None):
    ws = wb.create_sheet(name)
    if columns is None:
        columns = list(rows[0].keys()) if rows else []
    ws.append(columns)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    for row in rows:
        ws.append([flatten(row.get(column, '')) for column in columns])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for index, column in enumerate(columns, 1):
        width = min(max(len(str(column)) + 2, 12), 34)
        for values in ws.iter_rows(min_col=index, max_col=index, min_row=2):
            width = min(max(width, len(str(values[0].value or '')) + 2), 42)
        ws.column_dimensions[get_column_letter(index)].width = width
    return ws


def add_breakdown(ws, title, data, start_row):
    ws.cell(start_row, 1, title).fill = section_fill
    ws.cell(start_row, 1).font = Font(bold=True)
    ws.cell(start_row + 1, 1, 'Item').font = Font(bold=True)
    ws.cell(start_row + 1, 2, 'Count').font = Font(bold=True)
    for cell in ws[start_row + 1][0:2]:
        cell.fill = header_fill
        cell.font = header_font
    for offset, (key, value) in enumerate(sorted(data.items()), 2):
        ws.cell(start_row + offset, 1, key)
        ws.cell(start_row + offset, 2, value)
    return start_row + len(data) + 3


summary = wb.create_sheet('Summary')
summary['A1'] = 'GBMS Support App — Inception-to-Date Activity Report'
summary['A1'].font = title_font
summary.merge_cells('A1:D1')
summary['A3'] = 'Reporting period'
summary['B3'] = f"{report['inception']['firstRetainedRecord'][:10]} to {report['inception']['lastRetainedRecord'][:10]}"
summary['A4'] = 'Prepared'
summary['B4'] = report['generatedAt'][:10]
summary['A5'] = 'Source'
summary['B5'] = 'Firebase Firestore production activity records, project mofapp-60963'
summary['A7'] = 'Dataset'
summary['B7'] = 'All retained records'
summary['C7'] = 'Included operational records'
for cell in summary[7][0:3]:
    cell.fill = header_fill
    cell.font = header_font
included = report['included']
coverage_rows = [
    ('Users', report['coverage']['allUsers'], len(included['users'])),
    ('Tickets', report['coverage']['allTickets'], len(included['tickets'])),
    ('Ticket activities', report['coverage']['allTicketActivities'], len(included['activities'])),
    ('Notifications', report['coverage']['allNotifications'], len(included['notifications'])),
    ('System audit logs', report['coverage']['allAuditLogs'], len(included['auditLogs'])),
]
for row in coverage_rows:
    summary.append(row)
summary['A15'] = 'Test-ticket rule'
summary['B15'] = report['rules']['testTicketRule']
summary['A16'] = 'Excluded test tickets'
summary['B16'] = report['rules']['excludedTestTicketCount']
summary.column_dimensions['A'].width = 28
summary.column_dimensions['B'].width = 75
summary.column_dimensions['C'].width = 28
summary.freeze_panes = 'A8'

breakdowns = wb.create_sheet('Breakdowns')
row = 1
for title, key in [
    ('Users by role', 'usersByRole'),
    ('Users by institution type', 'usersByInstitutionType'),
    ('Users created by month', 'usersCreatedByMonth'),
    ('Tickets by month', 'ticketsByMonth'),
    ('Tickets by institution type', 'ticketsByInstitutionType'),
    ('Tickets by creator role', 'ticketsByCreatorRole'),
    ('Tickets by weekday', 'ticketsByDay'),
    ('Tickets by status', 'ticketsByStatus'),
    ('Tickets by priority', 'ticketsByPriority'),
    ('Tickets by category', 'ticketsByCategory'),
    ('Ticket activities by action', 'activitiesByAction'),
    ('Notifications by type', 'notificationsByType'),
    ('Audit events by action', 'auditByAction'),
    ('Audit events by month', 'auditByMonth'),
]:
    row = add_breakdown(breakdowns, title, report['breakdowns'][key], row)
breakdowns.column_dimensions['A'].width = 34
breakdowns.column_dimensions['B'].width = 16

add_table('Users', included['users'])
add_table('Tickets', included['tickets'])
all_ticket_rows = [
    {**ticket, 'reportInclusion': 'Included'}
    for ticket in included['tickets']
] + [
    {**ticket, 'reportInclusion': 'Excluded test ticket'}
    for ticket in report['excludedTestTickets']
]
add_table(f"All Tickets ({len(all_ticket_rows)})", all_ticket_rows)
add_table('Ticket Activities', included['activities'])
add_table('Notifications', included['notifications'])
add_table('Audit Logs', included['auditLogs'])
add_table('Excluded Tickets', report['excludedTestTickets'])

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    for row_cells in ws.iter_rows():
        for cell in row_cells:
            cell.alignment = Alignment(vertical='top', wrap_text=True)

output = ROOT / 'GBMS Activity Report - Updated v6.xlsx'
wb.save(output)
print(f'Wrote {output}')
