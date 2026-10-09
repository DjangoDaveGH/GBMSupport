import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
with (ROOT / 'september_2026_activity.json').open(encoding='utf-8') as f:
    report = json.load(f)

wb = Workbook()
summary = wb.active
summary.title = 'Summary'

navy = '0B3578'
gold = 'F2B705'
light = 'E6EFFB'

summary['A1'] = 'GBMS Support App — September 2026 Activity Report'
summary['A1'].font = Font(size=16, bold=True, color='FFFFFF')
summary['A1'].fill = PatternFill('solid', fgColor=navy)
summary.merge_cells('A1:D1')
summary['A3'] = 'Reporting period'
summary['B3'] = '1-30 September 2026 (full month)'
summary['A4'] = 'Prepared'
summary['B4'] = report['generatedAt'][:10]
summary['A5'] = 'Source'
summary['B5'] = 'Firebase Firestore production project mofapp-60963'
summary['A7'] = 'Metric'
summary['B7'] = 'Count'
for cell in summary[7]:
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor=navy)

metrics = [
    ('Users created', report['summary']['usersCreated']),
    ('New tickets', report['summary']['ticketsCreated']),
    ('Ticket activities', report['summary']['ticketActivities']),
    ('Notifications generated', report['summary']['notifications']),
    ('System audit-log events', report['summary']['auditLogs']),
]
for row, (label, value) in enumerate(metrics, 8):
    summary.cell(row, 1, label)
    summary.cell(row, 2, value)

summary['A15'] = 'Ticket status at extraction'
summary['A15'].font = Font(bold=True, color='FFFFFF')
summary['A15'].fill = PatternFill('solid', fgColor=gold)
for row, (label, value) in enumerate(report['breakdowns']['ticketsByStatusAtReportTime'].items(), 16):
    summary.cell(row, 1, label)
    summary.cell(row, 2, value)

def add_sheet(name, records):
    ws = wb.create_sheet(name)
    if not records:
        ws['A1'] = 'No records'
        return ws
    keys = list(records[0].keys())
    for col, key in enumerate(keys, 1):
        cell = ws.cell(1, col, key)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=navy)
        cell.alignment = Alignment(vertical='top')
    for row_index, record in enumerate(records, 2):
        for col, key in enumerate(keys, 1):
            value = record.get(key)
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            ws.cell(row_index, col, value)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col in range(1, len(keys) + 1):
        width = min(max(12, max(len(str(ws.cell(row, col).value or '')) for row in range(1, min(ws.max_row, 100) + 1)) + 2), 42)
        ws.column_dimensions[get_column_letter(col)].width = width
    return ws

users_created_for_report = []
for user in report['records']['usersCreated']:
    normalized_user = dict(user)
    if normalized_user.get('role') == 'mda_user':
        normalized_user['role'] = 'end_user'
    users_created_for_report.append(normalized_user)
add_sheet('Users Created', users_created_for_report)
add_sheet('Tickets Created', report['records']['ticketsCreated'])
add_sheet('Ticket Activity', report['records']['activities'])
add_sheet('Notifications', report['records']['notifications'])
add_sheet('Audit Logs', report['records']['auditLogs'])

def hours_between(start, end):
    if not start or not end:
        return ''
    from datetime import datetime
    start_dt = datetime.fromisoformat(start.replace('Z', '+00:00'))
    end_dt = datetime.fromisoformat(end.replace('Z', '+00:00'))
    return round((end_dt - start_dt).total_seconds() / 3600, 1)

ticket_detail_records = []
institution_type_by_id = {
    user.get('institutionId'): user.get('institutionType')
    for user in report['records']['usersCreated']
    if user.get('institutionId') and user.get('institutionType')
}
for institution_id in ['jaman-north-district', 'south-tongu-district', 'afadzato-south-district', 'north-tongu-district', 'gomoa-east-district', 'upper-east-assit-rba']:
    institution_type_by_id[institution_id] = 'MMDA'
institution_counts = {}
for ticket in sorted(report['records']['ticketsCreated'], key=lambda t: t.get('createdAt', '')):
    institution_id = ticket.get('institutionId') or '(not recorded)'
    institution_type = institution_type_by_id.get(ticket.get('institutionId'), 'Type not recorded')
    institution_counts[(institution_type, institution_id)] = institution_counts.get((institution_type, institution_id), 0) + 1
    ticket_detail_records.append({
        'ticketReference': ticket.get('ticketReference') or ticket.get('id'),
        'createdAt': ticket.get('createdAt'),
        'title': ticket.get('title'),
        'category': ticket.get('category'),
        'priority': ticket.get('priority'),
        'impact': ticket.get('impact'),
        'institutionType': institution_type,
        'institutionId': institution_id,
        'statusAtExtraction': ticket.get('status'),
        'assignedTo': ticket.get('assignedToName') or 'Unassigned',
        'firstRespondedAt': ticket.get('firstRespondedAt'),
        'resolvedAt': ticket.get('resolvedAt'),
        'responseHours': hours_between(ticket.get('createdAt'), ticket.get('firstRespondedAt')),
        'resolutionHours': hours_between(ticket.get('createdAt'), ticket.get('resolvedAt')),
        'resolutionNotes': ticket.get('resolutionNotes'),
    })
add_sheet('Ticket Detail', ticket_detail_records)
add_sheet('Institution Coverage', [
    {'institutionType': institution_type, 'institutionId': institution_id, 'tickets': count}
    for (institution_type, institution_id), count in sorted(institution_counts.items(), key=lambda item: (-item[1], item[0][1]))
])

response_hours = [x['responseHours'] for x in ticket_detail_records if x['responseHours'] != '']
resolution_hours = [x['resolutionHours'] for x in ticket_detail_records if x['resolutionHours'] != '']
add_sheet('Ticket Metrics', [
    {'metric': 'Tickets created in September', 'value': len(ticket_detail_records)},
    {'metric': 'Tickets with first response', 'value': len(response_hours)},
    {'metric': 'Average hours to first response', 'value': round(sum(response_hours) / len(response_hours), 1) if response_hours else 0},
    {'metric': 'Tickets with recorded resolution', 'value': len(resolution_hours)},
    {'metric': 'Average hours to resolution', 'value': round(sum(resolution_hours) / len(resolution_hours), 1) if resolution_hours else 0},
    {'metric': 'Distinct institutions represented', 'value': len({x['institutionId'] for x in ticket_detail_records if x['institutionId'] != '(not recorded)'})},
    {'metric': 'Tickets linked to MDAs', 'value': sum(x['tickets'] for x in [{'tickets': count, 'institutionType': institution_type} for (institution_type, _), count in institution_counts.items()] if x['institutionType'] == 'MDA')},
    {'metric': 'Tickets linked to MMDAs', 'value': sum(x['tickets'] for x in [{'tickets': count, 'institutionType': institution_type} for (institution_type, _), count in institution_counts.items()] if x['institutionType'] == 'MMDA')},
    {'metric': 'Tickets with institution type not recorded', 'value': sum(x['tickets'] for x in [{'tickets': count, 'institutionType': institution_type} for (institution_type, _), count in institution_counts.items()] if x['institutionType'] == 'Type not recorded')},
])

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)

output = ROOT / 'september_2026_activity_report.xlsx'
wb.save(output)
print(f'Wrote {output}')
