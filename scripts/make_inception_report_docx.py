import json
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

with open('scripts/inception_activity.json', encoding='utf-8') as f: r=json.load(f)
b=r['breakdowns']; inc=r['included']
doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
doc.styles['Normal'].font.name='Aptos'; doc.styles['Normal'].font.size=Pt(9)
t=doc.add_heading('GBMS Support App',0); t.alignment=WD_ALIGN_PARAGRAPH.CENTER
s=doc.add_paragraph('Inception-to-Date Activity Report'); s.alignment=WD_ALIGN_PARAGRAPH.CENTER; s.runs[0].bold=True
doc.add_paragraph(f"Reporting period: {r['inception']['firstRetainedRecord'][:10]} to {r['inception']['lastRetainedRecord'][:10]}\nPrepared: {r['generatedAt'][:10]}\nSource: Firebase Firestore production activity records, project mofapp-60963")
doc.add_heading('Executive summary',1)
doc.add_paragraph(f"The report covers all retained timestamped activity since the earliest available app record: {r['coverage']['allUsers']:,} user records, {len(inc['tickets']):,} operational tickets, {len(inc['activities']):,} related ticket activities, {len(inc['notifications']):,} related notifications, and {len(inc['auditLogs']):,} system audit-log events.")
doc.add_paragraph(f"Configured test-ticket rule: {r['rules']['testTicketRule']} This removed {r['rules']['excludedTestTicketCount']} test tickets and their related ticket activities and ticket-linked notifications.")
def table(title, data):
    doc.add_heading(title,2); tb=doc.add_table(rows=1,cols=2); tb.style='Light Shading Accent 1'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER; tb.rows[0].cells[0].text='Item'; tb.rows[0].cells[1].text='Count'
    for k,v in data.items(): c=tb.add_row().cells; c[0].text=str(k); c[1].text=f'{v:,}'
table('Overall coverage',{'Users':len(inc['users']),'Operational tickets':len(inc['tickets']),'Ticket activities':len(inc['activities']),'Notifications':len(inc['notifications']),'System audit logs':len(inc['auditLogs'])})
table('Users by role',b['usersByRole']); table('Users by institution type',b['usersByInstitutionType']); table('Users created by month',b['usersCreatedByMonth'])
table('Tickets by month',b['ticketsByMonth']); table('Tickets by institution type',b['ticketsByInstitutionType']); table('Tickets by weekday',b['ticketsByDay']); table('Tickets by status',b['ticketsByStatus']); table('Tickets by priority',b['ticketsByPriority']); table('Tickets by category',b['ticketsByCategory'])
table('Ticket activity actions',b['activitiesByAction']); table('Notifications by type',b['notificationsByType']); table('Audit actions',b['auditByAction']); table('Audit events by month',b['auditByMonth'])
doc.add_heading('Excluded test tickets',1); doc.add_paragraph(f"{r['rules']['excludedTestTicketCount']} tickets were excluded under the MMDA Tuesday/Friday rule. See excluded_test_tickets.csv for the full record.")
doc.add_heading('Data limitations',1); doc.add_paragraph('This is a report of retained Firestore records. Deleted or purged documents cannot be reconstructed from the live database. Presence heartbeats overwrite the current profile field and are not retained as historical sessions. Ordinary screen views and reads are not logged unless they create one of the covered records.')
doc.add_heading('Complete appendices',1); doc.add_paragraph('The accompanying CSV files contain the complete included ticket activity, audit logs, notifications, user-creation events, and excluded test-ticket list. The full structured export is inception_activity.json.')
doc.add_paragraph('Prepared for management submission.')
doc.save('scripts/inception_activity_report.docx'); print('Wrote inception_activity_report.docx')
